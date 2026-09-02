from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook

from backend.app.models import TransactionKind, TransactionType


EXPENSE_CATEGORIES = [
    "Продукты",
    "Транспорт",
    "Жилье",
    "Развлечения",
    "Здоровье",
    "Другое",
]
INCOME_CATEGORIES = ["Зарплата", "Фриланс", "Инвестиции", "Другое"]

EXPENSE_CATEGORY_MAP = {
    "супермаркеты": "Продукты",
    "фастфуд": "Продукты",
    "рестораны": "Продукты",
    "такси": "Транспорт",
    "местный транспорт": "Транспорт",
    "заправки": "Транспорт",
    "ж/д билеты": "Транспорт",
    "кино": "Развлечения",
    "цифровые товары": "Развлечения",
    "экосистема яндекс": "Развлечения",
    "аптеки": "Здоровье",
}
INCOME_CATEGORY_MAP = {
    "зарплата": "Зарплата",
    "проценты": "Инвестиции",
}
INCOME_BANK_CATEGORIES = {"зарплата", "проценты", "бонусы", "пополнения", "переводы"}
REQUIRED_HEADERS = {
    "Имя счёта",
    "Дата операции",
    "Сумма в валюте счёта",
    "Валюта счёта",
    "Статус",
    "Категория по-умолчанию",
    "Описание",
    "Учёт в аналитике",
}


class ImportFormatError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedTransaction:
    occurred_on: date
    amount_cents: int
    category: str
    comment: str | None
    source: str
    type: TransactionType
    kind: TransactionKind
    included_in_analytics: bool
    merchant: str | None
    bank_category: str | None
    mcc: str | None
    status: str
    fingerprint: str


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def merchant_key(value: str | None) -> str:
    return normalize_text(value).casefold()


def _money_to_cents(value: Any) -> int:
    try:
        amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as error:
        raise ImportFormatError(f"Некорректная сумма: {value}") from error
    return int(amount * 100)


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, pattern).date()
            except ValueError:
                continue
    raise ImportFormatError(f"Некорректная дата операции: {value}")


def _xlsx_rows(content: bytes) -> list[list[Any]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _xls_rows(content: bytes) -> list[list[Any]]:
    workbook = xlrd.open_workbook(file_contents=content)
    sheet = workbook.sheet_by_index(0)
    rows: list[list[Any]] = []
    for row_index in range(sheet.nrows):
        row: list[Any] = []
        for column_index in range(sheet.ncols):
            cell = sheet.cell(row_index, column_index)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
            else:
                row.append(cell.value)
        rows.append(row)
    return rows


def read_excel_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        rows = _xlsx_rows(content)
    elif suffix == ".xls":
        rows = _xls_rows(content)
    else:
        raise ImportFormatError("Поддерживаются только файлы .xlsx и .xls")

    if not rows:
        raise ImportFormatError("В файле нет данных")

    headers = [normalize_text(value) for value in rows[0]]
    missing = REQUIRED_HEADERS.difference(headers)
    if missing:
        raise ImportFormatError(
            "В файле отсутствуют обязательные столбцы: " + ", ".join(sorted(missing))
        )

    records = []
    for row in rows[1:]:
        if not any(value not in (None, "") for value in row):
            continue
        padded = row + [None] * max(0, len(headers) - len(row))
        records.append(dict(zip(headers, padded, strict=False)))
    return records


def _fingerprint(record: dict[str, Any], occurred_on: date, signed_cents: int) -> str:
    payload = {
        "account": normalize_text(record.get("Имя счёта")),
        "card": normalize_text(record.get("Номер карты")),
        "date": occurred_on.isoformat(),
        "raw_date": str(record.get("Дата операции")),
        "amount": signed_cents,
        "currency": normalize_text(record.get("Валюта счёта")),
        "status": normalize_text(record.get("Статус")),
        "bank_category": normalize_text(record.get("Категория по-умолчанию")),
        "description": normalize_text(record.get("Описание")),
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def parse_transaction(
    record: dict[str, Any],
    saved_rules: dict[tuple[str, TransactionType], str] | None = None,
) -> ParsedTransaction:
    saved_rules = saved_rules or {}
    occurred_on = _parse_date(record.get("Дата операции"))
    signed_cents = _money_to_cents(record.get("Сумма в валюте счёта"))
    if signed_cents == 0:
        raise ImportFormatError("Операция с нулевой суммой не поддерживается")

    status = normalize_text(record.get("Статус")) or "Ок"
    bank_category = normalize_text(record.get("Категория по-умолчанию")) or None
    bank_category_key = merchant_key(bank_category)
    description = normalize_text(record.get("Описание")) or None
    message = normalize_text(record.get("Сообщение")) or None
    account = normalize_text(record.get("Имя счёта"))
    source = account or "Не указан"
    analytics_enabled = merchant_key(record.get("Учёт в аналитике")) == "да"
    is_internal_transfer = merchant_key(description) == "между своими счетами"

    if status.casefold() != "ок":
        kind = TransactionKind.error
        transaction_type = TransactionType.expense if signed_cents < 0 else TransactionType.income
        included = False
    elif is_internal_transfer:
        kind = TransactionKind.internal_transfer
        transaction_type = TransactionType.expense if signed_cents < 0 else TransactionType.income
        included = False
    elif bank_category_key == "переводы":
        kind = TransactionKind.external_transfer
        transaction_type = TransactionType.expense if signed_cents < 0 else TransactionType.income
        included = analytics_enabled
    elif bank_category_key == "пополнения":
        kind = TransactionKind.top_up
        transaction_type = TransactionType.income
        included = analytics_enabled
    elif signed_cents > 0 and bank_category_key not in INCOME_BANK_CATEGORIES:
        kind = TransactionKind.refund
        transaction_type = TransactionType.expense
        included = analytics_enabled
    elif signed_cents > 0:
        kind = TransactionKind.income
        transaction_type = TransactionType.income
        included = analytics_enabled
    else:
        kind = TransactionKind.purchase
        transaction_type = TransactionType.expense
        included = analytics_enabled

    normalized_merchant = merchant_key(description)
    saved_category = saved_rules.get((normalized_merchant, transaction_type))
    if saved_category:
        category = saved_category
    elif transaction_type == TransactionType.expense:
        category = EXPENSE_CATEGORY_MAP.get(bank_category_key, "Другое")
    else:
        category = INCOME_CATEGORY_MAP.get(bank_category_key, "Другое")

    comment_parts = [part for part in (description, message) if part]
    comment = " — ".join(dict.fromkeys(comment_parts))[:255] or None
    mcc_value = record.get("MCC")
    mcc = normalize_text(mcc_value) or None

    return ParsedTransaction(
        occurred_on=occurred_on,
        amount_cents=abs(signed_cents),
        category=category,
        comment=comment,
        source=source[:100],
        type=transaction_type,
        kind=kind,
        included_in_analytics=included,
        merchant=description,
        bank_category=bank_category,
        mcc=mcc,
        status=status,
        fingerprint=_fingerprint(record, occurred_on, signed_cents),
    )
